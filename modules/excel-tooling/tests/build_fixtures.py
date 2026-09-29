"""Write the synthetic xlsx fixtures Calc opens in the libreoffice tests."""

from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def build() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    _titles(FIXTURES / "titles.xlsx")
    _epoch(FIXTURES / "epoch1904.xlsx")
    _empty(FIXTURES / "empty.xlsx")


def _titles(path: Path) -> None:
    wb = Workbook()
    empty = wb.active
    empty.title = "Empty"

    titles = wb.create_sheet("Titles")
    for column, name in enumerate(("Name", "City", "Amount", "When", "Active"), start=1):
        titles.cell(4, column, name)
    titles["A2"] = "Quarter report"
    titles["A5"] = "Ann"
    titles["B5"] = "Oslo"
    titles["C5"] = 2
    titles["D5"] = date(2024, 3, 15)
    titles["D5"].number_format = "yyyy-mm-dd"
    titles["E5"] = True
    titles["A7"] = "Bo"
    titles["C7"] = 3.5
    titles["D7"] = "=1+1"
    titles["E7"] = True

    types = wb.create_sheet("Types")
    headers = (
        "Text",
        "Whole",
        "Fraction",
        "Percent",
        "Money",
        "Date",
        "DateTime",
        "Time",
        "Flag",
        "Formula",
        "Err",
    )
    for column, name in enumerate(headers, start=1):
        types.cell(1, column, name)
    types["A2"] = "0012"
    types["A2"].number_format = "@"
    types["B2"] = 2
    types["C2"] = 2.5
    types["D2"] = 0.5
    types["D2"].number_format = "0%"
    types["E2"] = 1.5
    types["E2"].number_format = '"$"#,##0.00'
    types["F2"] = date(2024, 3, 15)
    types["F2"].number_format = "yyyy-mm-dd"
    types["G2"] = datetime(2024, 3, 15, 13, 45, 0)
    types["G2"].number_format = "yyyy-mm-dd hh:mm:ss"
    types["H2"] = time(13, 45, 0)
    types["H2"].number_format = "hh:mm:ss"
    types["I2"] = True
    types["J2"] = "=2+2"
    types["K2"] = "=NA()"

    gap = wb.create_sheet("Gap")
    gap["A1"] = "Name"
    gap["A2"] = "Keep"
    gap["A53"] = "Drop"

    dup = wb.create_sheet("Dup")
    dup["A1"] = "Name"
    dup["B1"] = "Name"

    wide = wb.create_sheet("Wide")
    wide["AA1"] = "Wide"

    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "Secret"
    hidden.sheet_state = "hidden"

    wb.save(path)


def _epoch(path: Path) -> None:
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    sheet = wb.active
    sheet.title = "Dates"
    sheet["A1"] = "Day"
    sheet["A2"] = date(2024, 3, 15)
    sheet["A2"].number_format = "yyyy-mm-dd"
    wb.save(path)


def _empty(path: Path) -> None:
    wb = Workbook()
    wb.active.title = "Empty"
    wb.save(path)


if __name__ == "__main__":
    build()
